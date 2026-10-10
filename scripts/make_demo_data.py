"""Generate a small, realistic demo library and (optionally) index it.

    python scripts/make_demo_data.py --index

Creates in data/raw/demo/:
  Q3_review_meeting.mp4  two people discussing a Q3 review over 4 slides (neural TTS voices)
  Q3_report.pdf          2-page report with a text layer
  Roadmap_2027.pptx      3-slide deck
  scanned_invoice.png    a "scanned" vendor invoice (no text layer - needs vision OCR / CLIP)

The facts cross-reference each other on purpose, so questions need several sources:
"Why did the Chennai launch slip and how much did we pay the vendor?"

Without internet (TTS unavailable) the meeting video is still created, silent, and its
transcript is built from the script so the demo works fully offline.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter, ImageFont  # noqa: E402

from voicevault.config import settings  # noqa: E402

VOICES = {"Priya": "en-IN-NeerjaNeural", "Rahul": "en-IN-PrabhatNeural"}
SR = 16000

SLIDES = [
    {"title": "Q3 Business Review", "bullets": ["Finance & Operations", "October 2026"],
     "caption": "Title slide: 'Q3 Business Review', Finance & Operations, October 2026."},
    {"title": "Revenue by Region (₹ Cr)", "chart": {"North": 4.2, "South": 3.1, "West": 3.8, "East": 2.4},
     "caption": "Slide 'Revenue by Region (₹ Cr)'. Bar chart: North 4.2, South 3.1, West 3.8, East 2.4. "
                "North is highest, East lowest."},
    {"title": "Risk: Chennai launch", "bullets": ["Moved from Q3 to Q4", "Vendor: Shakti Logistics",
                                                 "Warehouse fit-out delayed 6 weeks"],
     "caption": "Slide 'Risk: Chennai launch'. Bullets: moved from Q3 to Q4; vendor Shakti Logistics; "
                "warehouse fit-out delayed 6 weeks."},
    {"title": "Hiring plan Q4", "bullets": ["2 ML engineers - Bengaluru", "1 data analyst - Pune",
                                           "Budget approved: ₹48 L"],
     "caption": "Slide 'Hiring plan Q4'. Two ML engineers in Bengaluru, one data analyst in Pune, "
                "budget approved 48 lakh rupees."},
]

DIALOGUE = [  # (speaker, slide index, line)
    ("Priya", 0, "Okay everyone, let's start the Q3 business review."),
    ("Priya", 1, "Revenue first. North is our strongest region this quarter at four point two crore."),
    ("Rahul", 1, "South came in at three point one crore, which is below plan."),
    ("Rahul", 1, "Most of that gap is because the Chennai launch did not happen in Q3."),
    ("Priya", 2, "Right. The Chennai launch has slipped to Q4."),
    ("Rahul", 2, "Shakti Logistics delayed the warehouse fit-out by about six weeks."),
    ("Priya", 2, "Let's hold their final payment until the warehouse passes inspection."),
    ("Rahul", 2, "Agreed. I will send them a formal notice by Friday."),
    ("Priya", 3, "Last item is hiring. We are adding two ML engineers in Bengaluru."),
    ("Rahul", 3, "And one data analyst in Pune. The budget of forty eight lakh is approved."),
    ("Priya", 3, "Great. Action items: Rahul sends the vendor notice, I will open the job postings."),
]


# ------------------------------------------------------------------ images

def font(size: int):
    for name in ("DejaVuSans.ttf", "Arial.ttf", "arial.ttf", "/System/Library/Fonts/Helvetica.ttc"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def render_slide(spec: dict, path: Path) -> None:
    im = Image.new("RGB", (1280, 720), "#fbf8f2")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, 1280, 14], fill="#c8641e")
    d.text((80, 70), spec["title"], fill="#1d1c1a", font=font(54))
    if "chart" in spec:
        vals = spec["chart"]
        top = max(vals.values())
        for i, (k, v) in enumerate(vals.items()):
            x = 140 + i * 260
            h = int(400 * v / top)
            d.rectangle([x, 620 - h, x + 150, 620], fill="#c8641e" if i == 0 else "#7d7668")
            d.text((x + 30, 630), k, fill="#1d1c1a", font=font(30))
            d.text((x + 40, 620 - h - 44), f"{v}", fill="#1d1c1a", font=font(32))
    for i, b in enumerate(spec.get("bullets", [])):
        d.text((100, 220 + i * 90), "•  " + b, fill="#33302b", font=font(40))
    im.save(path)


def make_scan(path: Path) -> None:
    im = Image.new("L", (1000, 1300), 238)
    d = ImageDraw.Draw(im)
    lines = [("SHAKTI LOGISTICS PVT LTD", 44), ("Plot 17, Ambattur Industrial Estate, Chennai", 24),
             ("", 20), ("TAX INVOICE   No. INV-4471", 34), ("Date: 28 Sep 2026", 26),
             ("Bill to: VoiceVault Demo Co., Bengaluru", 26), ("", 20),
             ("Warehouse fit-out (phase 1)          1,80,000", 28),
             ("Racking installation                  65,000", 28), ("", 20),
             ("TOTAL (INR)                        2,45,000", 34), ("", 20),
             ("Payment terms: 30 days", 26), ("Status: FINAL PAYMENT ON HOLD", 26)]
    y = 80
    for text, size in lines:
        d.text((70, y), text, fill=40, font=font(size))
        y += size + 26
    d.ellipse([650, 1000, 900, 1180], outline=90, width=6)
    d.text((690, 1070), "RECEIVED", fill=90, font=font(34))
    rng = np.random.default_rng(3)
    noise = rng.normal(0, 9, (1300, 1000))
    im = Image.fromarray(np.clip(np.asarray(im, float) + noise, 0, 255).astype("uint8"))
    im = im.rotate(1.3, expand=True, fillcolor=210).filter(ImageFilter.GaussianBlur(0.6))
    im.convert("RGB").save(path)


# ------------------------------------------------------------- documents

def make_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    pages = [
        ["Q3 2026 Operations Report", "",
         "Revenue: total 13.5 Cr. North 4.2 Cr (best), West 3.8 Cr, South 3.1 Cr, East 2.4 Cr.",
         "South missed plan by 0.9 Cr, mainly due to the delayed Chennai launch.", "",
         "Customer support: average first response time fell from 9 h to 4 h after",
         "the new ticket routing went live in August."],
        ["Risks and mitigations", "",
         "1. Chennai launch moved to Q4. Vendor Shakti Logistics delayed the warehouse",
         "   fit-out by six weeks. Final payment withheld pending inspection.",
         "2. Two senior engineers resigned in September; backfill hiring approved.",
         "3. Cloud spend rose 18% quarter on quarter; reserved instances to be purchased.", "",
         "Travel policy reminder: reimbursement is capped at 5,000 rupees per trip."],
    ]
    for lines in pages:
        y = 780
        for i, line in enumerate(lines):
            c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 18 if i == 0 else 11)
            c.drawString(60, y, line)
            y -= 28 if i == 0 else 18
        c.showPage()
    c.save()


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    p = Presentation()
    slides = [("Roadmap 2027", ["Launch VoiceVault copilot for Hindi and English users",
                                "Expand Chennai warehouse to 2x capacity"]),
              ("Q1 2027 milestones", ["Hinglish speech model in production (target WER < 20%)",
                                      "Self-serve document upload for all teams"]),
              ("Open questions", ["Build vs buy for vendor management",
                                  "Data residency for client recordings"])]
    for title, bullets in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
    p.save(path)


# ----------------------------------------------------------------- meeting

async def _tts(text: str, voice: str, out: Path) -> None:
    import edge_tts

    await edge_tts.Communicate(text, voice).save(str(out))


def _to_pcm(mp3: Path) -> np.ndarray:
    raw = subprocess.run(["ffmpeg", "-loglevel", "error", "-i", str(mp3), "-f", "s16le",
                          "-ac", "1", "-ar", str(SR), "-"], capture_output=True, check=True).stdout
    return np.frombuffer(raw, np.int16)


def make_meeting(out_mp4: Path, work: Path) -> tuple[list[dict], bool]:
    """Returns (timed lines, real_audio)."""
    slide_pngs = []
    for i, spec in enumerate(SLIDES):
        p = work / f"slide_{i}.png"
        render_slide(spec, p)
        slide_pngs.append(p)

    real = True
    clips: list[np.ndarray] = []
    try:
        for i, (spk, _, line) in enumerate(DIALOGUE):
            mp3 = work / f"line_{i}.mp3"
            asyncio.run(_tts(line, VOICES[spk], mp3))
            clips.append(_to_pcm(mp3))
    except Exception as e:
        print(f"  ! neural TTS unavailable ({type(e).__name__}); meeting video will be silent "
              "and its transcript taken from the script")
        real = False
        clips = [np.zeros(int(SR * max(1.5, len(line.split()) / 2.6)), np.int16)
                 for _, _, line in DIALOGUE]

    gap = np.zeros(int(SR * 0.45), np.int16)
    audio, timed, t = [gap], [], len(gap) / SR
    for (spk, slide, line), clip in zip(DIALOGUE, clips, strict=True):
        dur = len(clip) / SR
        timed.append({"speaker": spk, "slide": slide, "text": line, "start": round(t, 2),
                      "end": round(t + dur, 2)})
        audio += [clip, gap]
        t += dur + len(gap) / SR
    wav = work / "meeting.wav"
    import wave
    with wave.open(str(wav), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(SR)
        w.writeframes(np.concatenate(audio).tobytes())

    # slide timeline: each slide shows from its first line until the next slide's first line
    starts = {}
    for ln in timed:
        starts.setdefault(ln["slide"], ln["start"] if ln["slide"] else 0.0)
    order = sorted(starts)
    concat = []
    for j, s in enumerate(order):
        end = starts[order[j + 1]] if j + 1 < len(order) else t
        concat += [f"file '{slide_pngs[s].as_posix()}'", f"duration {end - starts[s]:.2f}"]
    concat.append(f"file '{slide_pngs[order[-1]].as_posix()}'")
    (work / "slides.txt").write_text("\n".join(concat))
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-f", "concat", "-safe", "0",
                    "-i", str(work / "slides.txt"), "-i", str(wav), "-vf", "fps=10,format=yuv420p",
                    "-c:v", "libx264", "-tune", "stillimage", "-c:a", "aac", "-shortest", str(out_mp4)],
                   check=True)
    return timed, real


def scripted_ingest(video: Path, timed: list[dict]):
    """Offline stand-in for Whisper + diarization: transcript from the script, real keyframes."""
    from voicevault.ingest.chunker import build_chunks
    from voicevault.ingest.media import duration
    from voicevault.ingest.pipeline import _slug
    from voicevault.ingest.video import extract_keyframes
    from voicevault.schemas import IngestResult, Transcript, TranscriptSegment

    work = settings.processed_dir / _slug(video.stem)
    frames = extract_keyframes(video, work / "frames")
    for f in frames:  # we know what's on each slide, so caption it exactly
        slide = max((ln["slide"] for ln in timed if ln["start"] <= f.timestamp + 0.5), default=0)
        f.caption = SLIDES[slide]["caption"]
    tr = Transcript(source=video.name, language="en", duration=duration(video), segments=[
        TranscriptSegment(start=ln["start"], end=ln["end"], text=ln["text"], speaker=ln["speaker"])
        for ln in timed])
    res = IngestResult(source=video.name, modality="video", transcript=tr, keyframes=frames,
                       chunks=build_chunks(tr, frames, modality="video"), timings={})
    work.mkdir(parents=True, exist_ok=True)
    (work / "ingest.json").write_text(res.model_dump_json(indent=2), encoding="utf-8")
    return res


# -------------------------------------------------------------------- main

QUESTIONS = [
    "Why did the Chennai launch slip?",
    "How much do we owe Shakti Logistics and is it paid?",
    "Which region had the highest revenue? Show me the chart.",
    "What are the action items from the Q3 meeting?",
    "चेन्नई लॉन्च में देरी क्यों हुई?",
    "Hiring plan ka budget kitna hai?",
    "What is the travel reimbursement limit?",
]


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--out", type=Path, default=settings.raw_dir / "demo")
    p.add_argument("--index", action="store_true", help="also ingest + index everything")
    p.add_argument("--no-vision", action="store_true", help="skip vision OCR/captions when indexing")
    a = p.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)

    print("• documents")
    make_pdf(a.out / "Q3_report.pdf")
    make_pptx(a.out / "Roadmap_2027.pptx")
    make_scan(a.out / "scanned_invoice.png")
    print("• meeting video (text-to-speech)…")
    with tempfile.TemporaryDirectory() as tmp:
        video = a.out / "Q3_review_meeting.mp4"
        timed, real = make_meeting(video, Path(tmp))
    (a.out / "Q3_review_meeting.script.json").write_text(json.dumps(timed, indent=2))
    print(f"✓ demo files in {a.out}")

    if a.index:
        from voicevault.ingest.documents import ingest_document
        from voicevault.retrieval import Indexer, open_store

        store, temb, iemb = open_store()
        idx = Indexer(store, temb, iemb)
        for f in ("Q3_report.pdf", "Roadmap_2027.pptx", "scanned_invoice.png"):
            try:
                print("  ", idx.index(ingest_document(a.out / f, vision=not a.no_vision)))
            except Exception as e:
                print(f"   ✗ {f}: {e}")
        res = None
        if real:
            try:
                from voicevault.ingest.pipeline import ingest
                res = ingest(video, language="en", captions=not a.no_vision)
            except Exception as e:
                print(f"   ! Whisper pipeline unavailable ({e}); using scripted transcript")
        res = res or scripted_ingest(video, timed)
        print("  ", idx.index(res))
        store.close()

    print("\nTry asking:\n" + "\n".join(f"  • {q}" for q in QUESTIONS))


if __name__ == "__main__":
    main()
